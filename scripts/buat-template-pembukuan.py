"""Membuat assets/template-pembukuan-umkm.xlsx (hadiah formulir email, Fase 6 & 9 workflow).
Jalankan: python3 scripts/buat-template-pembukuan.py"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = 'Arial'
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
NORMAL = Font(name=F)
INPUT = Font(name=F, color='0000FF')
FILL_IN = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='D9EAD3')
RP = '"Rp"#,##0;-"Rp"#,##0;"-"'
THIN = Border(bottom=Side(style='thin', color='BBBBBB'))
KATEGORI = ['Penjualan', 'Pemasukan lain', 'Bahan baku', 'Kemasan', 'Gaji & upah', 'Sewa', 'Listrik, gas & air',
            'Promosi', 'Ongkos kirim', 'Pajak', 'Cicilan pinjaman', 'Peralatan', 'Lain-lain']
BARIS = 1000  # jumlah baris transaksi
BULAN = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni', 'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']

wb = Workbook()

# --- Petunjuk ---
ws = wb.active
ws.title = 'Petunjuk'
rows = [
    ('Template Pembukuan Usaha Kecil — LabaKita', TITLE),
    ('', None),
    ('Cara memakai', BOLD),
    ('1. Buka tab Pengaturan, isi tahun, jenis wajib pajak, dan saldo kas awal (sel kuning).', NORMAL),
    ('2. Catat setiap uang masuk dan keluar di tab Kas Harian. Pilih kategori dari daftar.', NORMAL),
    ('3. Tab Ringkasan menghitung otomatis omzet, pengeluaran, laba, dan perkiraan PPh Final UMKM 0,5% per bulan.', NORMAL),
    ('4. Tab HPP membantu menghitung HPP per produk dan harga jual dengan target margin.', NORMAL),
    ('', None),
    ('Keterangan warna', BOLD),
    ('Sel kuning dengan teks biru = diisi sendiri. Sel lain berisi rumus, jangan diubah.', NORMAL),
    ('', None),
    ('Catatan', BOLD),
    ('Baris pertama di Kas Harian adalah contoh. Hapus isinya (bukan barisnya) sebelum mulai mencatat.', NORMAL),
    ('Penjualan dicatat sebesar omzet kotor, sebelum dipotong biaya marketplace. Potongan dicatat sebagai pengeluaran.', NORMAL),
    ('PPh Final 0,5% mengikuti PP 55/2022 sebagaimana diubah PP 20/2026: orang pribadi tidak dikenai pajak atas omzet sampai Rp500 juta setahun.', NORMAL),
    ('Aturan pajak dapat berubah; cek pajak.go.id. Hasil adalah perkiraan, bukan nasihat pajak.', NORMAL),
    ('Kalkulator dan panduan lainnya: https://labakita.pages.dev', NORMAL),
]
for i, (text, font) in enumerate(rows, 1):
    c = ws.cell(i, 1, text)
    if font:
        c.font = font
ws.column_dimensions['A'].width = 110

# --- Pengaturan ---
ps = wb.create_sheet('Pengaturan')
ps['A1'] = 'Pengaturan'; ps['A1'].font = TITLE
items = [
    (3, 'Tahun pembukuan', 2026, '0', True),
    (4, 'Jenis wajib pajak', 'Orang pribadi', '@', True),
    (5, 'Saldo kas awal', 0, RP, True),
    (6, 'Omzet tidak kena pajak setahun', '=IF(B4="Orang pribadi",500000000,0)', RP, False),
    (7, 'Tarif PPh Final UMKM', 0.005, '0.0%', False),
]
for r, label, val, fmt, editable in items:
    ps.cell(r, 1, label).font = NORMAL
    c = ps.cell(r, 2, val)
    c.number_format = fmt
    c.font = INPUT if editable else NORMAL
    if editable:
        c.fill = FILL_IN
ps['C6'] = 'Rp500 juta hanya untuk orang pribadi; PT perorangan dan koperasi Rp0.'
ps['C7'] = 'PP 55/2022 jo. PP 20/2026. Cek perubahan aturan di pajak.go.id.'
ps['A9'] = 'Daftar kategori'; ps['A9'].font = BOLD
for i, k in enumerate(KATEGORI):
    ps.cell(10 + i, 1, k).font = NORMAL
ps.column_dimensions['A'].width = 32
ps.column_dimensions['B'].width = 20
ps.column_dimensions['C'].width = 60
dv_wp = DataValidation(type='list', formula1='"Orang pribadi,PT perorangan/koperasi"', allow_blank=False)
ps.add_data_validation(dv_wp); dv_wp.add('B4')
KAT_RANGE = f'Pengaturan!$A$10:$A${9 + len(KATEGORI)}'

# --- Kas Harian ---
ks = wb.create_sheet('Kas Harian')
ks['A1'] = 'Kas Harian'; ks['A1'].font = TITLE
ks['A2'] = 'Isi kolom A–E. Kolom Saldo dihitung otomatis.'; ks['A2'].font = NORMAL
heads = ['Tanggal', 'Keterangan', 'Kategori', 'Uang masuk', 'Uang keluar', 'Saldo']
for j, h in enumerate(heads, 1):
    c = ks.cell(4, j, h); c.font = BOLD; c.fill = HEAD
first, last = 5, 4 + BARIS
for r in range(first, last + 1):
    ks.cell(r, 1).number_format = 'DD/MM/YYYY'
    ks.cell(r, 4).number_format = RP
    ks.cell(r, 5).number_format = RP
    for j in range(1, 6):
        ks.cell(r, j).font = INPUT
    f = ks.cell(r, 6, f'=IF(AND(D{r}="",E{r}=""),"",Pengaturan!$B$5+SUM(D${first}:D{r})-SUM(E${first}:E{r}))')
    f.number_format = RP; f.font = NORMAL
# Contoh satu baris
import datetime
ks['A5'] = datetime.date(2026, 1, 2); ks['B5'] = 'Contoh: penjualan brownies 4 loyang'; ks['C5'] = 'Penjualan'; ks['D5'] = 620000
dv_k = DataValidation(type='list', formula1=KAT_RANGE, allow_blank=True)
ks.add_data_validation(dv_k); dv_k.add(f'C{first}:C{last}')
for col, w in zip('ABCDEF', [13, 42, 20, 16, 16, 18]):
    ks.column_dimensions[col].width = w
ks.freeze_panes = 'A5'

# --- Ringkasan ---
rs = wb.create_sheet('Ringkasan')
rs['A1'] = 'Ringkasan Bulanan'; rs['A1'].font = TITLE
rs['A2'] = '=CONCATENATE("Tahun ",Pengaturan!B3," · ",Pengaturan!B4)'; rs['A2'].font = NORMAL
heads = ['No', 'Bulan', 'Omzet (penjualan)', 'Pemasukan lain', 'Pengeluaran', 'Laba bersih', 'Omzet kumulatif', 'PPh Final 0,5%']
for j, h in enumerate(heads, 1):
    c = rs.cell(4, j, h); c.font = BOLD; c.fill = HEAD
D = f"'Kas Harian'!$A${first}:$A${last}"
C = f"'Kas Harian'!$C${first}:$C${last}"
IN = f"'Kas Harian'!$D${first}:$D${last}"
OUT = f"'Kas Harian'!$E${first}:$E${last}"
for i in range(12):
    r = 5 + i
    rs.cell(r, 1, i + 1).font = NORMAL
    rs.cell(r, 2, BULAN[i]).font = NORMAL
    period = f'{D},">="&DATE(Pengaturan!$B$3,A{r},1),{D},"<"&DATE(Pengaturan!$B$3,A{r}+1,1)'
    rs.cell(r, 3, f'=SUMIFS({IN},{C},"Penjualan",{period})')
    rs.cell(r, 4, f'=SUMIFS({IN},{period})-C{r}')
    rs.cell(r, 5, f'=SUMIFS({OUT},{period})')
    rs.cell(r, 6, f'=C{r}+D{r}-E{r}')
    rs.cell(r, 7, f'=SUM(C$5:C{r})')
    prev = '0' if r == 5 else f'G{r-1}'
    rs.cell(r, 8, f'=ROUND((MAX(0,G{r}-Pengaturan!$B$6)-MAX(0,{prev}-Pengaturan!$B$6))*Pengaturan!$B$7,0)')
    for j in range(3, 9):
        rs.cell(r, j).number_format = RP; rs.cell(r, j).font = NORMAL
rs['B17'].border = THIN
rs['B18'] = 'Total'; rs['B18'].font = BOLD
for j, col in zip(range(3, 9), 'CDEFGH'):
    if col == 'G':
        rs['G18'] = '=G16'
    else:
        rs.cell(18, j, f'=SUM({col}5:{col}16)')
    rs.cell(18, j).number_format = RP; rs.cell(18, j).font = BOLD
rs['A19'] = 'PPh dihitung dari omzet kategori "Penjualan" dan disetor paling lambat tanggal 15 bulan berikutnya.'
rs['A19'].font = Font(name=F, italic=True, color='666666')

rs['A21'] = 'Pengeluaran per kategori'; rs['A21'].font = BOLD
rs.cell(22, 2, 'Kategori').font = BOLD
for i in range(12):
    c = rs.cell(22, 3 + i, BULAN[i][:3]); c.font = BOLD; c.fill = HEAD
c = rs.cell(22, 15, 'Total'); c.font = BOLD; c.fill = HEAD
rs.cell(22, 2).fill = HEAD
from openpyxl.utils import get_column_letter as L
for k_i, k in enumerate(KATEGORI[2:]):
    r = 23 + k_i
    rs.cell(r, 2, k).font = NORMAL
    for m in range(12):
        col = 3 + m
        period = f'{D},">="&DATE(Pengaturan!$B$3,{m+1},1),{D},"<"&DATE(Pengaturan!$B$3,{m+2},1)'
        c = rs.cell(r, col, f'=SUMIFS({OUT},{C},$B{r},{period})')
        c.number_format = RP; c.font = NORMAL
    c = rs.cell(r, 15, f'=SUM(C{r}:N{r})'); c.number_format = RP; c.font = BOLD
rs.column_dimensions['A'].width = 5
rs.column_dimensions['B'].width = 20
for j in range(3, 16):
    rs.column_dimensions[L(j)].width = 16
rs.freeze_panes = 'C5'

# --- HPP ---
hs = wb.create_sheet('HPP')
hs['A1'] = 'Kalkulator HPP & Harga Jual'; hs['A1'].font = TITLE
inputs = [
    (3, 'Nama produk', 'Brownies loyang 22 cm', '@'),
    (4, 'Bahan baku per produksi', 300000, RP),
    (5, 'Tenaga kerja per produksi', 60000, RP),
    (6, 'Overhead per produksi (kemasan, gas, penyusutan)', 40000, RP),
    (7, 'Jumlah produk layak jual', 4, '#,##0'),
    (8, 'Potongan dari harga jual (marketplace/komisi/pajak)', 0, '0.0%'),
    (9, 'Target margin', 0.35, '0.0%'),
]
for r, label, val, fmt in inputs:
    hs.cell(r, 1, label).font = NORMAL
    c = hs.cell(r, 2, val); c.number_format = fmt; c.font = INPUT; c.fill = FILL_IN
outs = [
    (11, 'Total biaya produksi', '=B4+B5+B6', RP),
    (12, 'HPP per produk', '=IF(B7>0,B11/B7,0)', RP),
    (13, 'Harga jual minimal', '=IF(B8+B9<1,B12/(1-B8-B9),0)', RP),
    (14, 'Laba per produk setelah potongan', '=B13*(1-B8)-B12', RP),
]
for r, label, val, fmt in outs:
    hs.cell(r, 1, label).font = BOLD
    c = hs.cell(r, 2, val); c.number_format = fmt; c.font = BOLD
hs['A16'] = 'Contoh angka dari artikel "Cara Menghitung HPP Usaha Kue Rumahan" di LabaKita. Ganti sel kuning dengan angkamu.'
hs['A16'].font = Font(name=F, italic=True, color='666666')
hs.column_dimensions['A'].width = 52
hs.column_dimensions['B'].width = 24

wb.save('assets/template-pembukuan-umkm.xlsx')
print('assets/template-pembukuan-umkm.xlsx')
